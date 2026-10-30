"""Turning a slide deck into the PowerPoint file it's downloaded and
published as.

Decks use the institute template when `PPT_TEMPLATE_PATH` points at a
.pptx (its first layout for the title slide, its second, "Title and
Content", for the rest); otherwise a built-in 16:9 design in the logo's
colours.

Text never runs off a slide: each slide's bullets get the largest font
size that fits its body box (by an estimate of how the lines wrap), and a
slide too full even at the smallest size continues on "(cont.)" slides.
"""
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Emu, Inches, Pt

from app.config import settings

_BLUE = RGBColor(0x16, 0x84, 0xAC)
_DEEP = RGBColor(0x0D, 0x5F, 0x7E)
_INK = RGBColor(0x21, 0x1F, 0x1C)
_MUTED = RGBColor(0x5D, 0x6B, 0x70)

_WIDTH, _HEIGHT = Inches(13.333), Inches(7.5)
_BLANK_LAYOUT = 6  # in python-pptx's default template

# The built-in design's body box, and the font sizes tried, largest first.
_BODY_W, _BODY_H = Inches(11.8), Inches(4.95)
_BODY_SIZES = (24, 22, 20, 18)
# An average character is about half the font size wide in these fonts;
# a little extra keeps the estimate on the safe side.
_CHAR_WIDTH = 0.53
_LINE_HEIGHT = 1.2
_SPACE_AFTER = 0.6  # of the font size, after each bullet
_BULLET = "•  "


def _lines_needed(text: str, size: int, width: int) -> int:
    """How many lines `text` wraps to at `size` pt in a box `width` EMU wide."""
    per_line = max(1, int(Emu(width).pt / (size * _CHAR_WIDTH)))
    words, lines, used = text.split(), 1, 0
    for word in words:
        needed = len(word) if used == 0 else used + 1 + len(word)
        if needed <= per_line:
            used = needed
        else:
            lines += 1 + (len(word) - 1) // per_line  # an overlong word wraps on its own
            used = len(word) % per_line or per_line
    return lines


def _height(lines: list[str], size: int, width: int) -> float:
    """The points a list of paragraphs takes up at `size` pt."""
    return sum(_lines_needed(line, size, width) * size * _LINE_HEIGHT + size * _SPACE_AFTER for line in lines)


def fit_body(lines: list[str], width: int = _BODY_W, height: int = _BODY_H,
             sizes: tuple[int, ...] = _BODY_SIZES) -> list[tuple[list[str], int]]:
    """Splits a slide's bullets into the pages they need, each with the
    largest font size at which its bullets fit the body box: one page when
    they fit at the smallest size, more ("cont." slides) when they don't.
    """
    room = Emu(height).pt
    for size in sizes:
        if _height(lines, size, width) <= room:
            return [(lines, size)]
    smallest = sizes[-1]
    pages, page = [], []
    for line in lines:
        if page and _height([*page, line], smallest, width) > room:
            pages.append(page)
            page = []
        page.append(line)
    if page:
        pages.append(page)
    # Re-fit each page: a lighter page may get a bigger size back.
    return [next(((p, s) for s in sizes if _height(p, s, width) <= room), (p, smallest)) for p in pages]


def _title_size(title: str) -> int:
    """Long titles get smaller so they stay on at most two lines."""
    return 32 if len(title) <= 45 else 28 if len(title) <= 65 else 24


def sources_list(sources: list[dict]) -> list[str]:
    """"[n] Title" for each source, in citation order."""
    return [f"[{n}] {s['title']}" for n, s in enumerate(sources, start=1)]


def _text(frame, lines: list[str], size: int, color: RGBColor, bold: bool = False) -> None:
    frame.word_wrap = True
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(size * 0.6)
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def _builtin_slide(prs, title: str, lines: list[str], *, footer: str, number: int, size: int = 24,
                   cover: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.28), _HEIGHT)  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = _BLUE if not cover else _DEEP
    bar.line.fill.background()

    if cover:
        _text(slide.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11.3), Inches(1.6)).text_frame,
              [title], 44, _INK, bold=True)
        _text(slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(11.3), Inches(2)).text_frame,
              lines, 20, _MUTED)
        return slide

    _text(slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.8), Inches(1.2)).text_frame,
          [title], _title_size(title), _DEEP, bold=True)
    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), _BODY_W, _BODY_H).text_frame
    _text(body, [f"{_BULLET}{line}" for line in lines], size, _INK)
    _text(slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(10), Inches(0.4)).text_frame,
          [footer], 11, _MUTED)
    _text(slide.shapes.add_textbox(Inches(11.8), Inches(6.85), Inches(1), Inches(0.4)).text_frame,
          [str(number)], 11, _MUTED)
    return slide


def _template_body(prs) -> tuple[int, int]:
    """The width and height of the template's content placeholder."""
    layout = prs.slide_layouts[min(1, len(prs.slide_layouts) - 1)]
    body = next((p for p in layout.placeholders if p.placeholder_format.idx != 0), None)
    if body is None or not body.width or not body.height:
        return int(prs.slide_width * 0.85), int(prs.slide_height * 0.65)
    return body.width, body.height


def _template_slide(prs, title: str, lines: list[str], *, size: int | None = None, cover: bool = False):
    layout = prs.slide_layouts[0 if cover else min(1, len(prs.slide_layouts) - 1)]
    slide = prs.slides.add_slide(layout)
    placeholders = list(slide.placeholders)
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    body = next((p for p in placeholders if p.placeholder_format.idx != 0), None)
    if body is not None:
        frame = body.text_frame
        frame.word_wrap = True
        # Keeps the template's look; PowerPoint shrinks the text if it's still too much.
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            if size is not None:
                for run in paragraph.runs:
                    run.font.size = Pt(size)
    return slide


def deck_pptx(note: dict) -> bytes:
    """The deck as a .pptx: a title slide, one slide per outline entry (with
    its speaker notes), and a closing slide listing the sources.
    """
    template = settings.PPT_TEMPLATE_PATH.strip()
    if template:
        prs = Presentation(template)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = _WIDTH, _HEIGHT

    if template:
        # The bullet character and indent come from the template, so measure without ours.
        body_w, body_h = _template_body(prs)
        prefix = ""
    else:
        body_w, body_h = _BODY_W - Inches(0.2), _BODY_H  # textbox insets
        prefix = _BULLET

    number = 1

    def add(title: str, lines: list[str], notes: str = "", cover: bool = False) -> None:
        nonlocal number
        if cover:
            pages = [(lines, None)]
        else:
            pages = fit_body([f"{prefix}{line}" for line in lines], body_w, body_h)
            pages = [([line[len(prefix):] for line in page], size) for page, size in pages]
        for i, (page, size) in enumerate(pages):
            heading = title if i == 0 else f"{title} (cont.)"
            if template:
                slide = _template_slide(prs, heading, page, size=size, cover=cover)
            else:
                slide = _builtin_slide(prs, heading, page, footer=note["title"], number=number,
                                       size=size or 24, cover=cover)
            if notes and i == 0:
                slide.notes_slide.notes_text_frame.text = notes
            number += 1

    cover_lines = [line for line in (note.get("subject"), note["topic"]) if line]
    add(note["title"], cover_lines, cover=True)
    for s in note["slides"]:
        add(s["title"], s["bullets"], s.get("notes") or "")
    if note["sources"]:
        add("Sources", sources_list(note["sources"]))

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
